"""The profit page and its exports: who may open them, and what they show.

The definition of the figures is pinned in ``tests/unit/order/test_profit.py``;
this is the page around it: the permission (the store OWNER alone), the
range, bucket, comparison and breakdown, the CSV and the XLSX, and the
sidebar entry.
"""

from __future__ import annotations

import csv
import io
from datetime import UTC, date, datetime
from decimal import Decimal

import pytest
from django.test import Client
from django.urls import reverse
from django.utils import translation
from django.utils.translation import gettext
from djmoney.money import Money
from openpyxl import load_workbook

from admin.displays import money
from admin.reports import XLSX_CONTENT_TYPE
from order.enum.expense import ExpenseCategory
from order.enum.status import OrderStatus, PaymentStatus
from order.factories.expense import ExpenseFactory
from order.factories.order import OrderFactory
from product.factories import ProductFactory
from tenant.models import TenantMembershipRole
from tests.utils.staff import (
    bind_store_tenant,
    store_staff,
    store_tenant,
    unbind_store_tenant,
)
from tests.utils.turnover import snapshot_sale
from vat.factories import VatFactory

pytestmark = pytest.mark.django_db

D = Decimal

ROUTES = ["admin:profit", "admin:profit-csv", "admin:profit-xlsx"]
YEAR = {
    "start": "2026-01-01",
    "end": "2026-12-31",
    "bucket": "month",
    "compare": "",
    "group": "",
}


@pytest.fixture
def tenant():
    tenant = store_tenant("profit_page")
    previous = bind_store_tenant(tenant)
    yield tenant
    unbind_store_tenant(previous)


def _client(tenant, role=TenantMembershipRole.OWNER) -> Client:
    client = Client()
    client.force_login(
        store_staff(tenant, role=role),
        backend="tenant.auth_backends.PlatformStaffBackend",
    )
    return client


@pytest.fixture
def sale():
    """One order of 100.00 net goods costing 30.00 (124.00 at 24%) and 10.00
    shipping that left the store in March 2026, and a 25.00 carrier invoice
    and a 15.00 software bill the same month."""
    order = OrderFactory(
        num_order_items=0,
        status=OrderStatus.DELIVERED,
        payment_status=PaymentStatus.COMPLETED,
        shipping_price=Money(D("12.40"), "EUR"),
        payment_method_fee=Money(D("0"), "EUR"),
    )
    order.items.all().delete()
    order.items.create(
        product=ProductFactory(vat=VatFactory(value=D("24"))),
        price=Money(D("124.00"), "EUR"),
        quantity=1,
        vat_rate=D("24"),
        unit_cost=Money(D("30.00"), "EUR"),
    )
    ExpenseFactory(
        date=date(2026, 3, 5),
        category=ExpenseCategory.SHIPPING,
        amount_net=Money(D("25.00"), "EUR"),
    )
    ExpenseFactory(
        date=date(2026, 3, 6),
        category=ExpenseCategory.SOFTWARE,
        amount_net=Money(D("15.00"), "EUR"),
    )
    return snapshot_sale(order, datetime(2026, 3, 15, 12, tzinfo=UTC))


def _el(message):
    """``message`` in Greek: the language the admin is served in."""
    with translation.override("el"):
        return gettext(message)


def _cells(row):
    """A table row as the strings the page prints."""
    return [c["content"] if isinstance(c, dict) else c for c in row]


@pytest.mark.parametrize("name", ROUTES)
@pytest.mark.parametrize(
    ("role", "status"),
    [
        (TenantMembershipRole.STAFF, 403),
        (TenantMembershipRole.ADMIN, 403),
        (TenantMembershipRole.OWNER, 200),
    ],
)
def test_the_report_is_the_owners_alone(tenant, role, status, name):
    response = _client(tenant, role).get(reverse(name))

    assert response.status_code == status


def test_the_page_lists_the_months_with_the_whole_waterfall(tenant, sale):
    response = _client(tenant).get(reverse("admin:profit"), YEAR)
    # The admin is served in Greek, which writes 110,00.
    with translation.override("el"):
        expected = [
            money(D("110")),  # net revenue: 100 goods + 10 shipping
            money(D("30")),  # cost of goods sold
            money(D("80")),  # gross profit
            "72,7%",  # gross margin
            money(D("0")),  # payment fees
            money(D("25")),  # delivery costs
            money(D("55")),  # contribution
            money(D("15")),  # other expenses
            money(D("40")),  # operating profit
            "100,0%",  # cost coverage
            "—",  # no online payment, so no fee coverage to state
        ]

    assert response.status_code == 200
    table = response.context["table"]
    assert len(table["rows"]) == 13
    assert _cells(table["rows"][2])[1:12] == expected
    assert _cells(table["rows"][-1])[0] == _el("Total")
    assert _cells(table["rows"][-1])[1:12] == expected
    assert response.context["incomplete"] is False


def test_the_page_warns_when_a_cost_is_missing(tenant, sale):
    sale.items.update(unit_cost=None, unit_cost_currency=None)

    response = _client(tenant).get(reverse("admin:profit"), YEAR)

    assert response.context["incomplete"] is True
    [cost_coverage] = [
        kpi for kpi in response.context["kpis"] if kpi["icon"] == "fact_check"
    ]
    assert cost_coverage["tone"] == "warning"


def test_the_expenses_are_listed_by_category_biggest_first(tenant, sale):
    response = _client(tenant).get(reverse("admin:profit"), YEAR)

    rows = [_cells(row) for row in response.context["expense_table"]["rows"]]
    with translation.override("el"):
        assert rows == [
            [_el("Shipping (carrier invoices)"), money(D("25"))],
            [_el("Software and subscriptions"), money(D("15"))],
        ]


def test_the_comparison_adds_last_years_profit(tenant, sale):
    response = _client(tenant).get(
        reverse("admin:profit"), {**YEAR, "compare": "year"}
    )

    headers = [_cells([h])[0] for h in response.context["table"]["headers"]]
    assert _el("Operating profit last year") in headers
    assert _el("Change (%)") in headers


def test_a_breakdown_by_product_lists_its_goods_profit(tenant, sale):
    response = _client(tenant).get(
        reverse("admin:profit"), {**YEAR, "group": "product"}
    )

    [row] = response.context["group_table"]["rows"]
    cells = _cells(row)
    with translation.override("el"):
        assert cells[1:4] == [money(D("100")), money(D("30")), money(D("70"))]


def test_a_bad_range_is_refused_by_the_page_and_by_the_files(tenant):
    client = _client(tenant)
    bad = {**YEAR, "end": "2025-01-01"}

    assert client.get(reverse("admin:profit"), bad).context["form"].errors
    assert client.get(reverse("admin:profit-csv"), bad).status_code == 400
    assert client.get(reverse("admin:profit-xlsx"), bad).status_code == 400


def test_the_csv_has_plain_figures_and_a_total_row(tenant, sale):
    response = _client(tenant).get(reverse("admin:profit-csv"), YEAR)

    assert response["Content-Type"].startswith("text/csv")
    assert response.content.startswith("﻿".encode())
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert len(rows) == 14
    march = rows[3]
    assert march[0] == "2026-03"
    assert march[1:10] == [
        "110.00",
        "30.00",
        "80.00",
        "72.7",
        "0.00",
        "25.00",
        "55.00",
        "15.00",
        "40.00",
    ]
    assert rows[-1][0] == _el("Total")


def test_the_xlsx_has_a_sheet_of_numbers_and_one_per_breakdown(tenant, sale):
    response = _client(tenant).get(
        reverse("admin:profit-xlsx"), {**YEAR, "group": "product"}
    )

    assert response["Content-Type"] == XLSX_CONTENT_TYPE
    workbook = load_workbook(io.BytesIO(response.content))
    assert len(workbook.sheetnames) == 2
    periods = workbook.worksheets[0]
    march = [cell.value for cell in periods[4]]
    assert march[0] == "2026-03"
    assert march[1] == pytest.approx(110.0)
    assert march[9] == pytest.approx(40.0)


def test_the_sidebar_offers_it_to_the_owner_only(tenant):
    link = reverse("admin:profit")
    for role, offered in (
        (TenantMembershipRole.OWNER, True),
        (TenantMembershipRole.ADMIN, False),
        (TenantMembershipRole.STAFF, False),
    ):
        response = _client(tenant, role).get(reverse("admin:index"))

        assert (link in response.content.decode()) is offered
