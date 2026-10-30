"""The turnover page and its exports: who may open them, and what they show.

The definition of the figures is pinned in ``tests/unit/order/test_turnover.py``;
this is the page around it: the permission (store ADMIN and OWNER, not
STAFF), the range, bucket, comparison and breakdown, the CSV and the
XLSX, and the sidebar entry.
"""

from __future__ import annotations

import csv
import io
from datetime import UTC, datetime
from decimal import Decimal

import pytest
from django.test import Client
from django.urls import reverse
from django.utils import translation
from django.utils.translation import gettext
from djmoney.money import Money

from admin.displays import money
from admin.turnover import XLSX_CONTENT_TYPE
from order.enum.status import OrderStatus, PaymentStatus
from order.factories.order import OrderFactory
from product.factories import ProductFactory
from tenant.models import TenantMembershipRole
from tests.utils.staff import (
    bind_store_tenant,
    store_staff,
    store_tenant,
    unbind_store_tenant,
)
from tests.utils.turnover import snapshot_sale
from vat.factories import VatFactory

pytestmark = pytest.mark.django_db

D = Decimal


@pytest.fixture
def tenant():
    tenant = store_tenant("turnover_page")
    previous = bind_store_tenant(tenant)
    yield tenant
    unbind_store_tenant(previous)


def _client(tenant, role) -> Client:
    client = Client()
    client.force_login(
        store_staff(tenant, role=role),
        backend="tenant.auth_backends.PlatformStaffBackend",
    )
    return client


@pytest.fixture
def sale():
    """One order of 100.00 net goods (124.00 at 24%) and 10.00 shipping that
    left the store in March 2026."""
    order = OrderFactory(
        num_order_items=0,
        status=OrderStatus.DELIVERED,
        payment_status=PaymentStatus.COMPLETED,
        shipping_price=Money(D("10.00"), "EUR"),
        payment_method_fee=Money(D("0"), "EUR"),
    )
    order.items.all().delete()
    order.items.create(
        product=ProductFactory(vat=VatFactory(value=D("24"))),
        price=Money(D("124.00"), "EUR"),
        quantity=1,
        vat_rate=D("24"),
    )
    return snapshot_sale(order, datetime(2026, 3, 15, 12, tzinfo=UTC))


ROUTES = ["admin:turnover", "admin:turnover-csv", "admin:turnover-xlsx"]
YEAR = {
    "start": "2026-01-01",
    "end": "2026-12-31",
    "bucket": "month",
    "compare": "",
    "group": "",
}


def _el(message):
    """``message`` in Greek: the language the admin is served in."""
    with translation.override("el"):
        return gettext(message)


def _cells(row):
    """A table row as the strings the page prints."""
    return [c["content"] if isinstance(c, dict) else c for c in row]


@pytest.mark.parametrize("name", ROUTES)
@pytest.mark.parametrize(
    ("role", "status"),
    [
        (TenantMembershipRole.STAFF, 403),
        (TenantMembershipRole.ADMIN, 200),
        (TenantMembershipRole.OWNER, 200),
    ],
)
def test_the_report_needs_the_turnover_permission(tenant, role, status, name):
    response = _client(tenant, role).get(reverse(name))

    assert response.status_code == status


def test_the_page_lists_the_months_of_the_range_with_a_total(tenant, sale):
    response = _client(tenant, TenantMembershipRole.ADMIN).get(
        reverse("admin:turnover"), YEAR
    )
    # The admin is served in Greek, which writes 100,00.
    with translation.override("el"):
        expected = [
            money(D("100")),
            money(D("8.06")),
            money(D("0")),
            money(D("0")),
            money(D("108.06")),
            money(D("25.94")),
            money(D("134.00")),
            "1",
        ]

    assert response.status_code == 200
    table = response.context["table"]
    assert len(table["rows"]) == 13
    # Period, goods, shipping, fees, returns, turnover, VAT, gross, orders.
    assert _cells(table["rows"][2])[1:9] == expected
    total = _cells(table["rows"][-1])
    assert total[0] == _el("Total")
    assert total[1:9] == expected
    assert response.context["report"].totals.orders == 1


def test_the_page_defaults_to_this_year_by_month_against_last_year(
    tenant, sale
):
    response = _client(tenant, TenantMembershipRole.ADMIN).get(
        reverse("admin:turnover")
    )

    form = response.context["form"]
    assert form.is_valid()
    assert form.cleaned_data["bucket"] == "month"
    assert form.compare_year
    assert form.cleaned_data["start"].month == 1


def test_a_comparison_adds_the_year_before_to_the_table(tenant, sale):
    response = _client(tenant, TenantMembershipRole.ADMIN).get(
        reverse("admin:turnover"), {**YEAR, "compare": "year"}
    )

    headers = _cells(response.context["table"]["headers"])
    assert _el("Net turnover last year") in headers
    assert _el("Change (%)") in headers


def test_a_breakdown_lists_the_groups(tenant, sale):
    response = _client(tenant, TenantMembershipRole.ADMIN).get(
        reverse("admin:turnover"), {**YEAR, "group": "vat_rate"}
    )

    assert response.status_code == 200
    assert response.context["group_table"]["rows"][0][0] == "24%"


def test_the_year_shortcuts_link_to_their_year(tenant, sale):
    response = _client(tenant, TenantMembershipRole.ADMIN).get(
        reverse("admin:turnover"), YEAR
    )

    years = response.context["years"]
    assert years[-1]["label"] == "2026"
    assert "start=2026-01-01" in years[-1]["href"]


@pytest.mark.parametrize(
    "bad",
    [
        {"start": "abc"},
        {"end": "2026-02-30"},
        {"start": "2026-12-31", "end": "2026-01-01"},
        {"bucket": "decade"},
        {"group": "colour"},
        {"compare": "never"},
    ],
)
@pytest.mark.parametrize("name", ["admin:turnover-csv", "admin:turnover-xlsx"])
def test_an_export_refuses_what_the_form_refuses(tenant, sale, name, bad):
    response = _client(tenant, TenantMembershipRole.ADMIN).get(
        reverse(name), {**YEAR, **bad}
    )

    assert response.status_code == 400


def test_the_page_shows_a_form_error_instead_of_a_table(tenant, sale):
    response = _client(tenant, TenantMembershipRole.ADMIN).get(
        reverse("admin:turnover"), {**YEAR, "end": "2025-01-01"}
    )

    assert response.status_code == 200
    assert response.context["form"].errors["end"]
    assert "report" not in response.context


def test_the_page_caps_its_rows_but_the_export_does_not(tenant, sale):
    wide = {**YEAR, "start": "2020-01-01", "bucket": "day"}
    client = _client(tenant, TenantMembershipRole.ADMIN)

    page = client.get(reverse("admin:turnover"), wide)
    export = client.get(reverse("admin:turnover-csv"), wide)

    assert page.context["form"].errors["bucket"]
    assert export.status_code == 200


def test_the_csv_holds_the_periods_and_the_total(tenant, sale):
    response = _client(tenant, TenantMembershipRole.OWNER).get(
        reverse("admin:turnover-csv"), YEAR
    )

    assert response.status_code == 200
    assert response["Content-Type"].startswith("text/csv")
    assert (
        'filename="turnover-2026-01-01-2026-12-31.csv"'
        in response["Content-Disposition"]
    )
    assert response.content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert len(rows) == 14
    assert rows[0][:2] == [_el("Period"), _el("Goods")]
    assert rows[3][:9] == [
        "2026-03",
        "100.00",
        "8.06",
        "0.00",
        "0.00",
        "108.06",
        "25.94",
        "134.00",
        "1",
    ]
    assert rows[-1][0] == _el("Total")
    assert rows[-1][5] == "108.06"


def test_the_csv_of_a_breakdown_is_one_row_per_period_and_group(tenant, sale):
    response = _client(tenant, TenantMembershipRole.OWNER).get(
        reverse("admin:turnover-csv"), {**YEAR, "group": "vat_rate"}
    )

    rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
    assert rows[0][:2] == [_el("Period"), _el("VAT rate")]
    assert rows[1][:3] == ["2026-03", "24%", "108.06"]
    assert len(rows) == 2


def test_the_xlsx_holds_numbers_not_text(tenant, sale):
    from openpyxl import load_workbook

    response = _client(tenant, TenantMembershipRole.OWNER).get(
        reverse("admin:turnover-xlsx"), {**YEAR, "group": "vat_rate"}
    )

    assert response.status_code == 200
    assert response["Content-Type"] == XLSX_CONTENT_TYPE
    assert response["Content-Disposition"].endswith('.xlsx"')
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == [_el("Turnover"), _el("VAT rate")]
    sheet = book[_el("Turnover")]
    march = [cell.value for cell in sheet[4]]
    assert march[0] == "2026-03"
    assert march[1] == 100
    assert march[5] == pytest.approx(108.06)
    assert sheet.cell(row=4, column=2).number_format == "#,##0.00"
    assert [c.value for c in book[_el("VAT rate")][2]][:3] == [
        "2026-03",
        "24%",
        pytest.approx(108.06),
    ]


def test_the_sidebar_offers_the_report_to_those_who_may_open_it(tenant):
    for role, offered in (
        (TenantMembershipRole.ADMIN, True),
        (TenantMembershipRole.STAFF, False),
    ):
        response = _client(tenant, role).get(reverse("admin:index"))
        link = reverse("admin:turnover")

        assert (link in response.content.decode()) is offered
